#!/usr/bin/env python3
"""Gate de validación estructural para un XLSForm antes de subirlo a Kobo.

Uso:
    python3 validate_xlsform.py encuesta.xlsx

Exit codes: 0 = sin errores estructurales, 1 = errores (no subir), 2 = error de uso.
Si `pyxform` está instalado, además intenta compilar el form a XForm de verdad
(el mismo validador que usa Kobo/ODK) — best-effort, no bloquea si falla el import.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Falta openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(2)

REF_RE = re.compile(r"\$\{([a-zA-Z0-9_]+)\}")


def load_sheet(wb, name):
    if name not in wb.sheetnames:
        return None, []
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return [], []
    header = [str(h or "").strip() for h in rows[0]]
    data = [dict(zip(header, r)) for r in rows[1:] if any(c is not None for c in r)]
    return header, data


def main():
    if len(sys.argv) != 2:
        print(__doc__, file=sys.stderr)
        sys.exit(2)
    path = Path(sys.argv[1])
    if not path.exists():
        print(f"No existe: {path}", file=sys.stderr)
        sys.exit(2)

    wb = openpyxl.load_workbook(path, data_only=True)
    survey_header, survey = load_sheet(wb, "survey")
    choices_header, choices = load_sheet(wb, "choices")

    errors = []
    warnings = []

    if survey_header is None:
        errors.append("Falta la hoja 'survey'.")
    if choices_header is None:
        errors.append("Falta la hoja 'choices'.")
    if errors:
        _report(errors, warnings)
        sys.exit(1)

    label_col = next((c for c in survey_header if c.startswith("label")), None)

    lists_available = {}
    for c in choices:
        ln = (c.get("list_name") or "").strip()
        if not ln:
            continue
        lists_available.setdefault(ln, 0)
        lists_available[ln] += 1

    names_seen = set()
    group_stack = []
    prior_names_at_row = []  # nombres visibles ANTES de cada fila (para chequear referencias)

    for i, row in enumerate(survey, start=2):
        prior_names_at_row.append(set(names_seen))
        rtype = (row.get("type") or "").strip()
        name = (row.get("name") or "").strip()

        if rtype == "begin_group":
            if not name:
                errors.append(f"Fila {i}: begin_group sin 'name'.")
            group_stack.append((name, i))
            continue
        if rtype == "end_group":
            if not group_stack:
                errors.append(f"Fila {i}: end_group sin begin_group correspondiente.")
            else:
                open_name, open_row = group_stack.pop()
                if name and name != open_name:
                    errors.append(
                        f"Fila {i}: end_group '{name}' no matchea el begin_group "
                        f"'{open_name}' abierto en la fila {open_row}."
                    )
            continue

        if not rtype:
            errors.append(f"Fila {i}: falta 'type'.")
            continue

        if rtype not in ("start", "end", "today", "note") and not name:
            errors.append(f"Fila {i} ({rtype}): falta 'name'.")

        if name:
            if name in names_seen:
                errors.append(f"Fila {i}: 'name'={name!r} repetido (ya aparece antes).")
            names_seen.add(name)

        if rtype.startswith("select_one ") or rtype.startswith("select_multiple "):
            list_name = rtype.split(" ", 1)[1].strip()
            if list_name not in lists_available:
                errors.append(
                    f"Fila {i} ({name}): referencia la lista '{list_name}' que no "
                    f"existe en 'choices' (o no tiene ninguna opción)."
                )

        for col in ("relevant", "constraint", "calculation"):
            expr = row.get(col) or ""
            for ref in REF_RE.findall(str(expr)):
                if ref not in prior_names_at_row[-1]:
                    errors.append(
                        f"Fila {i} ({name or rtype}), columna '{col}': referencia "
                        f"'${{{ref}}}' que no existe o aparece después en el formulario."
                    )

        if label_col:
            label = str(row.get(label_col) or "")
            blob = label.lower()
            for signal in ("nombre", "apellido", "documento", "cédula", "cedula", "dni",
                           "teléfono", "telefono", "celular", "whatsapp", "correo",
                           "email", "dirección", "direccion", "fecha de nacimiento"):
                if signal in blob:
                    warnings.append(f"Fila {i} ({name}): posible PII — {label!r} (señal: {signal!r})")
                    break

    if group_stack:
        for open_name, open_row in group_stack:
            errors.append(f"begin_group '{open_name}' (fila {open_row}) nunca se cerró con end_group.")

    _report(errors, warnings)

    if not errors:
        _try_pyxform(path)

    sys.exit(1 if errors else 0)


def _report(errors, warnings):
    print("=" * 72)
    print("  validate_xlsform.py — gate estructural para XLSForm")
    print("=" * 72)
    if warnings:
        print(f"\n⚠ Avisos ({len(warnings)}) — no bloquean, pero revisá antes de publicar:")
        for w in warnings:
            print(f"  - {w}")
    if errors:
        print(f"\n⛔ Errores ({len(errors)}) — NO subir a Kobo:")
        for e in errors:
            print(f"  - {e}")
        print("\n⛔ RESULTADO: hay errores estructurales.")
    else:
        print("\n✓ RESULTADO: sin errores estructurales.")
    print("=" * 72)


def _try_pyxform(path: Path):
    try:
        from pyxform.xls2json import parse_file_to_json
        from pyxform.builder import create_survey_element_from_dict
    except ImportError:
        print("\n(pyxform no está instalado — salteando compilación real a XForm. "
              "`pip install pyxform` para el chequeo completo tipo Kobo/ODK.)")
        return
    try:
        json_survey = parse_file_to_json(str(path))
        survey = create_survey_element_from_dict(json_survey)
        # validate=False: compila a XForm sin pasar por ODK Validate (requiere
        # Java). Igual atrapa errores de estructura/columnas de pyxform.
        survey.to_xml(validate=False)
        print("\n✓ pyxform: el form compila correctamente a XForm.")
    except Exception as e:  # pyxform levanta excepciones variadas según el error
        print(f"\n⛔ pyxform: el form NO compila — {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
