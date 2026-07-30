#!/usr/bin/env python3
"""Construye un XLSForm (survey/choices/settings) desde el CSV/XLSX intermedio.

Uso:
    python3 build_xlsform.py borrador.csv --out encuesta.xlsx --titulo "Encuesta AMA"
    python3 build_xlsform.py borrador.xlsx --out encuesta.xlsx --aplanado

Ver references/esquema_intermedio.md para las columnas esperadas del CSV/XLSX
de entrada, references/matrices_y_grupos.md para el patrón de matriz, y
references/logica_de_salto.md para la sintaxis de `condicion`.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import re
import sys
import unicodedata
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Falta openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(2)

REQUIRED_COL = "pregunta"
KNOWN_COLS = [
    "seccion", "matriz", "pregunta", "tipo", "opciones", "lista", "ayuda",
    "requerido", "condicion", "nombre", "validacion",
]

SURVEY_HEADER = [
    "type", "name", "label::Spanish (es)", "hint::Spanish (es)", "required",
    "appearance", "constraint", "constraint_message", "relevant", "default",
    "calculation",
]
CHOICES_HEADER = ["list_name", "name", "label::Spanish (es)"]

VALIDACIONES = {
    "telefono:co": dict(type="integer",
                         constraint=". >= 3000000000 and . <= 3599999999",
                         constraint_message="Tu número de teléfono debe tener 10 dígitos."),
    "telefono:bo": dict(type="integer",
                         constraint="regex(., '^[0-9]{8}$')",
                         constraint_message="Tú número de telefono debe tener solo 8 dígitos."),
    "documento:co": dict(type="integer",
                          constraint="regex(., '^[0-9]{6,10}$')",
                          constraint_message="Tu número de identificación debe tener entre 6 y 10 dígitos."),
    "documento:bo": dict(type="integer",
                          constraint="regex(., '^[0-9]{7,8}$')",
                          constraint_message="Tu número de identificación debe tener 7 u 8 digitos."),
}

PII_SIGNALS = [
    "nombre", "apellido", "documento", "cedula", "dni", "telefono",
    "celular", "whatsapp", "correo", "email", "direccion",
    "fecha de nacimiento",
]


class BuildError(Exception):
    pass


def slug(text: str, maxlen: int = 40) -> str:
    text = unicodedata.normalize("NFKD", text or "")
    text = "".join(c for c in text if not unicodedata.combining(c))
    text = text.lower()
    text = re.sub(r"[^a-z0-9]+", "_", text).strip("_")
    text = re.sub(r"_+", "_", text)
    if not text:
        text = "campo"
    if text[0].isdigit():
        text = f"c_{text}"
    words = text.split("_")
    out = ""
    for w in words:
        cand = (out + "_" + w).strip("_") if out else w
        if len(cand) > maxlen:
            break
        out = cand
    return out or text[:maxlen]


def dedupe_name(name: str, used: set) -> str:
    if name not in used:
        used.add(name)
        return name
    i = 2
    while f"{name}_{i}" in used:
        i += 1
    final = f"{name}_{i}"
    used.add(final)
    return final


def strip_accents_lower(s: str) -> str:
    s = unicodedata.normalize("NFKD", s or "")
    return "".join(c for c in s if not unicodedata.combining(c)).lower()


def read_rows(path: Path) -> list[dict]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        header = [strip_accents_lower(str(h or "").strip()) for h in next(rows_iter)]
        rows = []
        for r in rows_iter:
            if all(c is None for c in r):
                continue
            rows.append({header[i]: ("" if r[i] is None else str(r[i]).strip())
                         for i in range(len(header)) if i < len(r)})
        return rows
    with open(path, newline="", encoding="utf-8-sig") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(f, dialect=dialect)
        rows = []
        for r in reader:
            norm = {strip_accents_lower(k.strip()): (v or "").strip() for k, v in r.items() if k}
            rows.append(norm)
        return rows


def parse_options(raw: str) -> list[str]:
    return [o.strip() for o in raw.split("|") if o.strip()]


def infer_tipo(row: dict) -> str:
    opciones = row.get("opciones", "")
    pregunta = row.get("pregunta", "")
    ayuda = row.get("ayuda", "")
    blob = f"{pregunta} {ayuda}".lower()
    if opciones:
        if any(k in blob for k in ("selecciona todas", "marca todas", "seleccionar todas", "(selecciona maximo", "(selecciona máximo".lower())):
            return "multiple"
        return "unica"
    if "fecha de nacimiento" in blob:
        return "fecha"
    if "edad" in blob:
        return "numero"
    if not pregunta.strip():
        raise BuildError(f"Fila sin 'pregunta' y sin tipo=nota explícito: {row}")
    return "texto"


class Builder:
    def __init__(self, aplanado: bool):
        self.aplanado = aplanado
        self.survey: list[dict] = []
        self.choices: list[dict] = []
        self.used_names: set = set()
        self.list_cache: dict = {}  # tuple(opciones) -> list_name
        self.used_list_names: set = set()
        self.cur_seccion = None
        self.cur_matriz = None
        self.cur_matriz_list = None
        self.cur_matriz_grouped = False

    def resolve_name(self, row: dict, fallback_text: str) -> str:
        explicit = row.get("nombre", "")
        base = slug(explicit) if explicit else slug(fallback_text)
        return dedupe_name(base, self.used_names)

    def resolve_list(self, row: dict, opciones: list[str]) -> str:
        explicit = row.get("lista", "")
        if explicit:
            list_name = slug(explicit, maxlen=20)
        else:
            key = tuple(opciones)
            if key in self.list_cache:
                return self.list_cache[key]
            h = hashlib.sha1("|".join(opciones).encode("utf-8")).hexdigest()[:6]
            list_name = f"l{h}"
        if list_name not in self.used_list_names:
            self.used_list_names.add(list_name)
            used_choice_names = set()
            for opt in opciones:
                cname = dedupe_name(slug(opt, maxlen=30), used_choice_names)
                self.choices.append({"list_name": list_name, "name": cname, "label::Spanish (es)": opt})
        self.list_cache[tuple(opciones)] = list_name
        return list_name

    def parse_condicion(self, cond: str) -> str:
        cond = (cond or "").strip()
        if not cond:
            return ""
        if cond.lower().startswith("xpath:"):
            return cond[len("xpath:"):].strip()
        neg = "!=" in cond
        sep = "!=" if neg else "="
        if sep not in cond:
            raise BuildError(
                f"No entiendo la condición '{cond}'. Usá 'nombre=valor', "
                f"'nombre!=valor' o 'xpath:<expresión>'. Ver references/logica_de_salto.md"
            )
        name, valores = cond.split(sep, 1)
        name = name.strip()
        if name not in self.used_names:
            raise BuildError(
                f"La condición '{cond}' referencia '{name}', que no existe "
                f"(o todavía no aparece) en el formulario. Revisá el 'nombre' "
                f"de la pregunta ancla."
            )
        valores = [slug(v, maxlen=30) for v in valores.split("|") if v.strip()]
        op = "!=" if neg else "="
        return " or ".join(f"${{{name}}} {op} '{v}'" for v in valores)

    def close_matriz(self):
        if self.cur_matriz is not None:
            if self.cur_matriz_grouped:
                self.survey.append({"type": "end_group", "name": self.cur_matriz})
            self.cur_matriz = None
            self.cur_matriz_list = None
            self.cur_matriz_grouped = False

    def close_seccion(self):
        self.close_matriz()
        if self.cur_seccion is not None:
            self.survey.append({"type": "end_group", "name": self.cur_seccion})
            self.cur_seccion = None

    def maybe_open_groups(self, row: dict, opciones_for_matriz: list[str]):
        seccion_raw = row.get("seccion", "")
        matriz_raw = row.get("matriz", "")
        seccion_name = slug(seccion_raw) if seccion_raw else None
        matriz_name = slug(matriz_raw) if matriz_raw else None

        if seccion_name != self.cur_seccion:
            self.close_seccion()
            self.cur_seccion = seccion_name
            if seccion_name:
                self.survey.append({
                    "type": "begin_group", "name": seccion_name,
                    "label::Spanish (es)": seccion_raw,
                })

        if matriz_name != self.cur_matriz:
            self.close_matriz()
            if matriz_name:
                self.cur_matriz = matriz_name
                list_name = self.resolve_list(row, opciones_for_matriz)
                self.cur_matriz_list = list_name
                if self.aplanado:
                    # Sin field-list ni header: cada ítem queda como select_one
                    # suelto (comparten lista igual). La condición NO se
                    # consume acá — se aplica fila por fila más abajo.
                    self.cur_matriz_grouped = False
                    return False
                self.cur_matriz_grouped = True
                relevant = self.parse_condicion(row.get("condicion", ""))
                self.survey.append({
                    "type": "begin_group", "name": matriz_name,
                    "appearance": "field-list", "relevant": relevant,
                })
                self.survey.append({
                    "type": f"select_one {list_name}",
                    "name": f"{matriz_name}_header",
                    "label::Spanish (es)": " ",
                    "appearance": "label",
                    "required": "no",
                })
                return True  # condicion ya consumida en el begin_group
        return False

    def add_row(self, row: dict, titulo_falta_pregunta_ok=False):
        tipo = row.get("tipo", "").strip().lower() or infer_tipo(row)
        pregunta = row.get("pregunta", "")
        ayuda = row.get("ayuda", "")
        requerido = row.get("requerido", "").strip().lower()
        condicion_raw = row.get("condicion", "")
        validacion = row.get("validacion", "").strip().lower()
        opciones = parse_options(row.get("opciones", "")) if tipo in ("unica", "multiple") else []

        cond_consumida = False
        if tipo in ("unica", "multiple") and row.get("matriz", ""):
            cond_consumida = self.maybe_open_groups(row, opciones)
        else:
            self.maybe_open_groups(row, opciones)

        if tipo == "nota":
            name = self.resolve_name(row, pregunta or "nota")
            self.survey.append({
                "type": "note", "name": name,
                "label::Spanish (es)": pregunta, "hint::Spanish (es)": ayuda,
                "required": "no",
            })
            return

        name = self.resolve_name(row, pregunta)
        entry = {"name": name, "label::Spanish (es)": pregunta, "hint::Spanish (es)": ayuda}
        entry["required"] = "no" if requerido == "no" else "yes"

        if tipo in ("unica", "multiple"):
            if not opciones:
                raise BuildError(f"Pregunta '{pregunta}' es {tipo} pero no trae 'opciones'.")
            list_name = self.cur_matriz_list if row.get("matriz", "") else self.resolve_list(row, opciones)
            prefix = "select_one" if tipo == "unica" else "select_multiple"
            entry["type"] = f"{prefix} {list_name}"
            if row.get("matriz", "") and not self.aplanado:
                entry["appearance"] = "list-nolabel"
        elif tipo == "texto":
            entry["type"] = "text"
        elif tipo == "numero":
            entry["type"] = "integer"
        elif tipo == "fecha":
            entry["type"] = "date"
        else:
            raise BuildError(f"Tipo desconocido '{tipo}' en pregunta '{pregunta}'. "
                              f"Usá texto/numero/fecha/unica/multiple/nota.")

        if validacion:
            if validacion.startswith("xpath:"):
                entry["constraint"] = validacion[len("xpath:"):].strip()
            elif validacion.startswith("edad:") and "-" in validacion:
                lo, hi = validacion[len("edad:"):].split("-", 1)
                entry["type"] = "integer"
                entry["constraint"] = f". >= {int(lo)} and . <= {int(hi)}"
                entry["constraint_message"] = f"La edad debe estar entre {int(lo)} y {int(hi)} años."
            elif validacion in VALIDACIONES:
                v = VALIDACIONES[validacion]
                entry["type"] = v["type"]
                entry["constraint"] = v["constraint"]
                entry["constraint_message"] = v["constraint_message"]
            else:
                raise BuildError(f"Atajo de validación desconocido '{validacion}' en '{pregunta}'.")

        if condicion_raw and not cond_consumida:
            entry["relevant"] = self.parse_condicion(condicion_raw)

        self.used_names.add(name)
        self.survey.append(entry)

    def finish(self):
        self.close_seccion()

    def pii_warnings(self) -> list[str]:
        warnings = []
        for row in self.survey:
            label = row.get("label::Spanish (es)", "") or ""
            blob = strip_accents_lower(label)
            for signal in PII_SIGNALS:
                if signal in blob:
                    warnings.append(f"{row.get('name')!r}: {label!r} (señal: {signal!r})")
                    break
        return warnings


def write_xlsform(builder: "Builder", out_path: Path, titulo: str, idioma: str):
    wb = openpyxl.Workbook()
    ws_survey = wb.active
    ws_survey.title = "survey"
    ws_survey.append(SURVEY_HEADER)
    for entry in builder.survey:
        ws_survey.append([entry.get(col, "") for col in SURVEY_HEADER])

    ws_choices = wb.create_sheet("choices")
    ws_choices.append(CHOICES_HEADER)
    for c in builder.choices:
        ws_choices.append([c.get(col, "") for col in CHOICES_HEADER])

    ws_settings = wb.create_sheet("settings")
    ws_settings.append(["form_title", "style", "default_language"])
    ws_settings.append([titulo, "pages", idioma])

    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("entrada", type=Path, help="CSV o XLSX intermedio (ver references/esquema_intermedio.md)")
    ap.add_argument("--out", type=Path, required=True, help="Ruta del XLSForm de salida (.xlsx)")
    ap.add_argument("--titulo", default="Encuesta AMA", help="form_title en la hoja settings")
    ap.add_argument("--idioma", default="Spanish (es)", help="default_language en la hoja settings")
    ap.add_argument("--aplanado", action="store_true",
                     help="No arma matrices field-list: cada ítem queda como select_one suelto "
                          "(para canales sin grilla, ej. bot/Typeform). Ver references/matrices_y_grupos.md")
    args = ap.parse_args()

    if not args.entrada.exists():
        print(f"No existe: {args.entrada}", file=sys.stderr)
        sys.exit(2)

    rows = read_rows(args.entrada)
    if not rows:
        print("El archivo de entrada no tiene filas.", file=sys.stderr)
        sys.exit(2)

    builder = Builder(aplanado=args.aplanado)
    try:
        for i, row in enumerate(rows, start=2):  # fila 2 = primera fila de datos (1 = header)
            try:
                builder.add_row(row)
            except BuildError as e:
                raise BuildError(f"Fila {i}: {e}") from e
        builder.finish()
    except BuildError as e:
        print(f"⛔ {e}", file=sys.stderr)
        sys.exit(1)

    write_xlsform(builder, args.out, args.titulo, args.idioma)
    print(f"✓ {args.out} — {len(builder.survey)} filas de survey, "
          f"{len(builder.choices)} opciones en {len(builder.used_list_names)} listas.")

    warnings = builder.pii_warnings()
    if warnings:
        print("\n⚠ Preguntas de identificación (PII) detectadas — anonimizar en la ingesta:")
        for w in warnings:
            print(f"  - {w}")


if __name__ == "__main__":
    main()
